from openpyxl import load_workbook
from rest_framework import status
from rest_framework.parsers import MultiPartParser
from rest_framework.response import Response
from rest_framework.views import APIView

from .repositories import DjangoORMClientRepository
from .serializers import ClientSerializer
from .services import ClientService


class ClientListView(APIView):
    service = ClientService(repo=DjangoORMClientRepository())

    def get(self, request):
        clients = self.service.list_clients()
        serializer = ClientSerializer(clients, many=True)
        return Response(serializer.data)


class ClientDetailView(APIView):
    service = ClientService(repo=DjangoORMClientRepository())

    def get(self, request, customer_id: int):
        client = self.service.get_client(customer_id=customer_id)
        serializer = ClientSerializer(client)
        return Response(serializer.data)

    def put(self, request, customer_id: int):
        current = self.service.get_client(customer_id=customer_id)
        serializer = ClientSerializer(current, data=request.data, partial=True)
        serializer.is_valid(raise_exception=True)
        data = serializer.validated_data

        updated = self.service.update_client(
            customer_id=customer_id,
            name=data.get("name", current.name),
            email=data.get("email", current.email),
            country=data.get("country", current.country),
            age=data.get("age", current.age),
        )
        out = ClientSerializer(updated)
        return Response(out.data)

    def delete(self, request, customer_id: int):
        self.service.delete_client(customer_id=customer_id)
        return Response(status=status.HTTP_204_NO_CONTENT)


class ClientImportView(APIView):
    parser_classes = [MultiPartParser]
    service = ClientService(repo=DjangoORMClientRepository())

    def post(self, request):
        file_obj = request.FILES.get("file")
        if file_obj is None:
            return Response(
                {"detail": "Missing file. Send multipart/form-data with field 'file'."},
                status=status.HTTP_400_BAD_REQUEST,
            )

        try:
            wb = load_workbook(filename=file_obj, data_only=True)
        except Exception:
            return Response({"detail": "Invalid Excel file"}, status=status.HTTP_400_BAD_REQUEST)

        if "Clientes" not in wb.sheetnames:
            return Response(
                {"detail": "Sheet 'Clientes' not found"},
                status=status.HTTP_400_BAD_REQUEST,
            )

        sheet = wb["Clientes"]
        header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
        columns = [str(v).strip() if v is not None else "" for v in header]

        rows = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value in (None, "") for value in row):
                continue
            rows.append(
                {
                    "customer_id": row[0] if len(row) > 0 else None,
                    "name": row[1] if len(row) > 1 else None,
                    "email": row[2] if len(row) > 2 else None,
                    "country": row[3] if len(row) > 3 else None,
                    "age": row[4] if len(row) > 4 else None,
                }
            )

        result = self.service.import_clients(rows=rows, columns=columns)
        return Response(result)

