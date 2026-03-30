from io import BytesIO

from django.db import IntegrityError
from django.test import TestCase
from openpyxl import Workbook
from rest_framework.test import APIClient

from .models import Client
from .services import ClientService


def build_excel(rows: list[list[object]]) -> BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "Clientes"
    ws.append(["customer_id", "name", "email", "country", "age"])
    for row in rows:
        ws.append(row)
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    output.name = "clientes.xlsx"
    return output


class ClientApiTests(TestCase):
    def setUp(self):
        self.api = APIClient()

    def test_import_clients_returns_summary_and_errors(self):
        payload = build_excel(
            [
                [1, "Alice", "alice@example.com", "AR", 20],
                [2, "", "invalid-email", "AR", 17],
            ]
        )
        response = self.api.post("/clients/import", {"file": payload}, format="multipart")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["summary"]["total_records"], 2)
        self.assertEqual(response.data["summary"]["inserted"], 1)
        self.assertEqual(response.data["summary"]["errors"], 1)
        self.assertEqual(Client.objects.count(), 1)

    def test_crud_clients(self):
        Client.objects.create(
            customer_id=10,
            name="Bob",
            email="bob@example.com",
            country="UY",
            age=30,
        )

        list_response = self.api.get("/clients")
        self.assertEqual(list_response.status_code, 200)
        self.assertEqual(len(list_response.data), 1)

        get_response = self.api.get("/clients/10")
        self.assertEqual(get_response.status_code, 200)
        self.assertEqual(get_response.data["name"], "Bob")

        put_response = self.api.put(
            "/clients/10",
            {
                "customer_id": 10,
                "name": "Bobby",
                "email": "bobby@example.com",
                "country": "UY",
                "age": 31,
            },
            format="json",
        )
        self.assertEqual(put_response.status_code, 200)
        self.assertEqual(put_response.data["name"], "Bobby")

        delete_response = self.api.delete("/clients/10")
        self.assertEqual(delete_response.status_code, 204)
        self.assertEqual(delete_response.content, b"")
        self.assertEqual(Client.objects.count(), 0)

    def test_import_clients_missing_file_returns_400(self):
        response = self.api.post("/clients/import", {}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertIn("Missing file", response.data["detail"])

    def test_import_clients_invalid_excel_returns_400(self):
        payload = BytesIO(b"not-an-excel-file")
        payload.name = "clientes.xlsx"

        response = self.api.post("/clients/import", {"file": payload}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["detail"], "Invalid Excel file")

    def test_import_clients_missing_sheet_returns_400(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "OtherSheet"
        ws.append(["customer_id", "name", "email", "country", "age"])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        output.name = "clientes.xlsx"

        response = self.api.post("/clients/import", {"file": output}, format="multipart")

        self.assertEqual(response.status_code, 400)
        self.assertEqual(response.data["detail"], "Sheet 'Clientes' not found")

    def test_import_clients_invalid_columns_returns_error_details(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Clientes"
        ws.append(["id", "name", "email", "country", "age"])
        ws.append([1, "Alice", "alice@example.com", "AR", 20])
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        output.name = "clientes.xlsx"

        response = self.api.post("/clients/import", {"file": output}, format="multipart")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["summary"]["inserted"], 0)
        self.assertEqual(response.data["summary"]["errors"], 1)
        self.assertIn("Invalid columns", response.data["error_details"][0]["errors"][0])

    def test_client_detail_returns_404_for_missing_customer(self):
        get_response = self.api.get("/clients/999")
        put_response = self.api.put(
            "/clients/999",
            {"name": "Nobody", "email": "nobody@example.com", "country": "AR", "age": 25},
            format="json",
        )
        delete_response = self.api.delete("/clients/999")

        self.assertEqual(get_response.status_code, 404)
        self.assertEqual(put_response.status_code, 404)
        self.assertEqual(delete_response.status_code, 404)

    def test_import_handles_integrity_error_and_reports_row_error(self):
        class RepoStub:
            def customer_id_exists(self, *, customer_id: int) -> bool:
                return False

            def create_client(self, **kwargs):
                raise IntegrityError("duplicate key")

        service = ClientService(repo=RepoStub())
        result = service.import_clients(
            rows=[{"customer_id": 1, "name": "Alice", "email": "alice@example.com", "country": "AR", "age": 20}],
            columns=["customer_id", "name", "email", "country", "age"],
        )

        self.assertEqual(result["summary"]["total_records"], 1)
        self.assertEqual(result["summary"]["inserted"], 0)
        self.assertEqual(result["summary"]["errors"], 1)
        self.assertEqual(result["error_details"][0]["customer_id"], 1)
        self.assertIn("customer_id already exists in database", result["error_details"][0]["errors"])
